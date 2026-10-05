#!/usr/bin/env python3
"""Parse the QA test-case Excel workbooks into canonical JSON documents for MongoDB.

Reads every test sheet, cleans it with simple rules, and writes three document
types that all go into one MongoDB collection (qa_knowledge):

  test_case  one test row (the full row is in content_md) + runs, links
  context    shared setup: parent test-data blocks, preconditions, configuration
  feature    one per module / interop feature (from the IntropsFeatures sheet)

Each document keeps only what the LLM and the search tools need:
  content_md         what the LLM reads (and what gets embedded)
  runs               status / owner per track, for filters and counts
  hierarchy          module / group / section, for filters
  entities, environments, relations   for joins and links
  source             sheet + row, for citations
  embedding          placeholder; vector is null until the embedding step runs

Version 1 is rules only. These are intentionally NOT filled yet:
  - embedding vectors (field exists, vector = null)
  - DEPENDS_ON relations
  - free-text entities (screen names, buttons)

Usage:
  python parse_testcases.py
  python parse_testcases.py --input "ALL MODULES_TCs.xlsx" --out output

Outputs (written to --out):
  qa_knowledge.jsonl   every document, one per line (for mongoimport)
  test_cases.json      test_case documents (pretty, for reading)
  contexts.json        context documents
  features.json        feature documents
  review_report.md     counts per sheet + rows a person should check
"""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  python -m pip install openpyxl")

SCHEMA_VERSION = "1.0"


# ---------------------------------------------------------------------------
# 1. CONFIGURATION
#    Edit these tables when a sheet changes. The code below should not need
#    changes for normal layout tweaks.
# ---------------------------------------------------------------------------

# Workbook file name (without .xlsx) -> (ID code, display domain)
WORKBOOKS = {
    "ALL MODULES_TCs": ("CORE", "Core"),
    "InteropsTeamFunctionalityDetails": ("INTEROP", "Interop"),
}

# Header text (lower case, single spaces) -> field role
HEADER_SYNONYMS = {
    "#": "tc_number", "tc#": "tc_number", "s no": "tc_number", "s.no": "tc_number",
    "s.n": "tc_number", "sno": "tc_number",
    "test scenario": "title", "scenario": "title", "test name": "title",
    "testcasename": "title", "testcase name": "title", "testcase names": "title",
    "test nameid": "label",
    "steps": "steps", "test steps": "steps", "test step": "steps", "testcase steps": "steps",
    "expected results": "expected", "expected result": "expected", "expected": "expected",
    "excepted result": "expected", "excepted results": "expected", "expected name": "expected",
    "results": "expected", "verification": "expected", "verification 1": "expected",
    "verification 2": "expected_2", "verification 3": "expected_3",
    "test data": "test_data", "testdata": "test_data",
    "preconditions": "preconditions", "preconditons": "preconditions",
    "precondition": "preconditions",
    "note": "notes", "notes": "notes",
    "highlevel": "high_level", "high level": "high_level",
    "tc source": "tc_source",
    "status": "status", "comment": "comment", "comments": "comment",
    "reg mfm": "status:Reg MFM",
}
# Headers that are typos of a real name (flagged as HEADER_TYPO_FIXED)
TYPO_HEADERS = {"excepted result", "excepted results", "preconditons", "expected name"}

# Execution tracks: the two databases in ALL MODULES
TRACK_SCOPES = {"Firebird": "HIGH_LEVEL", "MySQL": "FULL_REGRESSION"}

# Per-sheet rules the general logic cannot guess.
#   mode:              "tests" (default) | "skip" | "features" | "reference"
#   columns:           column letter -> role, overrides / fills the header row
#   first_data_row:    sheet has no header row; rows before this are setup notes
#   title_only_rows:   "section" (default) or "test" - what a row with only a title is
#   section_resets_group: a standalone label row ends the current test-data group
SHEET_CONFIG = {
    # ALL MODULES workbook
    "Overview": {"mode": "skip"},   # summary counts; compute from runs[] instead
    "Split": {"mode": "skip"},      # team assignment table
    "Nurse": {"columns": {"A": "title", "C": "high_level"}},   # header row has no A / C labels
    "Provider": {"columns": {"A": "title", "B": "steps"}},     # header row has no A / B labels
    # Interop workbook
    "IntropsFeatures": {"mode": "features"},
    "eRx": {"columns": {"F": "status:default", "G": "comment:default"},
            "section_resets_group": True},
    "eLabs": {"columns": {"G": "status:col_G"}},
    "KN02.Direct": {"columns": {"A": "tc_number", "B": "title", "C": "steps", "D": "expected"},
                    "first_data_row": 4},
    "PMX+": {"title_only_rows": "test"},
    "Eligibility Testing": {"title_only_rows": "test"},
    "Fhir_appointment_updated": {"mode": "reference"},
}

# IntropsFeatures "Introps List" name -> the sheet that holds its tests
FEATURE_SHEET_MAP = {
    "DTI": "DTI", "DRP": "DRP", "IIS testing": "IIS", "eLabs": "eLabs", "eFax": "eFax",
    "eRx": "eRx", "ePA": "ePA", "Carequlity": "Carequality", "ADT": "HUB_ADT&SIU",
    "DFT": "HUB_DFT", "837": "HUB_837", "EMR": "EMR", "HIMMS(VXU\\QBP)": "HUB_HIMMS",
}

STATUS_MAP = {
    "pass": "PASS", "passed": "PASS", "fail": "FAIL", "failed": "FAIL",
    "n/a": "NA", "na": "NA", "not applicable": "NA",
    "blocker": "BLOCKED", "blocked": "BLOCKED",
    "not run": "NOT_RUN", "pending": "PENDING", "in progress": "IN_PROGRESS",
}
CONFLICT_COMMENT_RE = re.compile(r"pending|fail|issue|bug|not working|error", re.I)

# Parent test-data block keys (eRx "Testcase - N" rows) -> field name
KV_KEYS = {
    "patient name": "patient_name", "patient id": "patient_id", "pharmacy": "pharmacy",
    "primary diagnosis": "primary_dx", "secondary diagnosis": "secondary_dx",
    "substitutions": "substitutions", "medication": "medication", "drug description": "medication",
    "quantity": "quantity", "potency unit code": "unit", "refills": "refills",
}
SNAPSHOT_KEYS = ["patient_id", "patient_name", "medication", "pharmacy", "primary_dx", "substance_class"]


def _p(etype, value, regex, standard=None, case=False):
    return {"type": etype, "value": value, "standard": standard,
            "re": re.compile(regex, 0 if case else re.I)}


# Dictionary + regex entities. value=None keeps the matched text.
ENTITY_PATTERNS = [
    # NCPDP SCRIPT (e-prescribing)
    _p("message_type", "CancelRxResponse",
       r"\bRx\s?Can(?:ce|e)l\s?Response\b|\bCan(?:ce|e)l\s?(?:Rx\s?)?Response\b", "NCPDP SCRIPT"),
    _p("message_type", "CancelRx", r"\bCan(?:ce|e)l\s?Rx\b(?!\s?Response)", "NCPDP SCRIPT"),
    _p("message_type", "NewRx", r"\bNew\s?Rx\b", "NCPDP SCRIPT"),
    _p("message_type", "RxRenewalRequest", r"\b(?:Rx\s?)?Renewal\s?Request\b", "NCPDP SCRIPT"),
    _p("message_type", "RxRenewalResponse", r"\b(?:Rx\s?)?Renewal\s?Response\b", "NCPDP SCRIPT"),
    _p("message_type", "RxChangeRequest", r"\bRx\s?Change\s?Request\b", "NCPDP SCRIPT"),
    _p("message_type", "RxChangeResponse", r"\bRx\s?Change\s?Response\b", "NCPDP SCRIPT"),
    _p("message_type", "RxFill", r"\bRx\s?Fill\b", "NCPDP SCRIPT"),
    _p("message_type", "PAInitiationRequest", r"\bPA\s?In\w*?tion\s?Request\b", "NCPDP SCRIPT"),
    _p("message_type", "PAInitiationResponse", r"\bPA\s?In\w*?tion\s?Response\b", "NCPDP SCRIPT"),
    _p("message_type", "PAAppealRequest", r"\bPA\s?Appeal\s?Request\b|\bAppeal\s?request\b", "NCPDP SCRIPT"),
    _p("message_type", "PAAppealResponse", r"\bPA\s?Appeal\s?Res(?:ponse)?\b", "NCPDP SCRIPT"),
    _p("message_type", "PACancelRequest", r"\bPA\s?Cancel\s?Req(?:uest)?\b", "NCPDP SCRIPT"),
    _p("message_type", "PARequest", r"\bPA\s?Request\b", "NCPDP SCRIPT"),
    _p("message_type", "PAResponse", r"\bPA\s?Response\b", "NCPDP SCRIPT"),
    # HL7 v2
    _p("message_type", "ADT_A04", r"\bADT[\s_^-]?A04\b|\bA04\b", "HL7 v2"),
    _p("message_type", "ADT_A08", r"\bADT[\s_^-]?A08\b|\bA08\b", "HL7 v2"),
    _p("message_type", "SIU_S12", r"\bSIU[\s_^-]?S12\b|\bS12\b", "HL7 v2"),
    _p("message_type", "SIU_S14", r"\bSIU[\s_^-]?S14\b|\bS14\b", "HL7 v2"),
    _p("message_type", "SIU_S15", r"\bSIU[\s_^-]?S15\b|\bS15\b", "HL7 v2"),
    _p("message_type", "ORU_R01", r"\bORU(?:\^R01)?\b", "HL7 v2"),
    _p("message_type", "VXU", r"\bVXU\b", "HL7 v2"),
    _p("message_type", "QBP", r"\bQBP\b", "HL7 v2"),
    _p("message_type", "RSP", r"\bRSP\b", "HL7 v2"),
    _p("message_type", "DFT", r"\bDFT\b", "HL7 v2"),
    # X12, CDA, FHIR, Direct
    _p("message_type", "837", r"\b837\b", "X12"),
    _p("message_type", "270/271", r"\b270\s*[,/&]\s*#?271\b", "X12"),
    _p("message_type", "CDA", r"\bC-?CDA\b|\bCDA\b", "HL7 CDA"),
    _p("message_type", "FHIR", r"\bFHIR\b", "HL7 FHIR"),
    _p("message_type", "Direct Message", r"\bDirect\s?message", "Direct"),
    # Systems
    _p("system", "Office Practicum (OP)", r"\bOP\b|\bop application\b", case=True),
    _p("system", "Surescripts", r"\bSure\s?scripts?\b"),
    _p("system", "Platform Manager", r"\bPlatform\s?Manager\b"),
    _p("system", "Kno2", r"\bKno2(?:fy)?\b"),
    _p("system", "Concord", r"\bConc[ao]rd\b"),
    _p("system", "DrFirst", r"\bDr\.?\s?First\b"),
    _p("system", "Arcadia", r"\bArcadia\b"),
    _p("system", "Carequality", r"\bCarequ?a?lity\b"),
    _p("system", "EMR Direct", r"\bEMR\s?Direct\b|\bemrdirect\b"),
    _p("system", "Quest Diagnostics", r"\bQuest\b"),
    _p("system", "IIS (immunization registry)", r"\bIIS\b"),
    _p("system", "CLC tool", r"\bCLC\b"),
    _p("system", "Crabkey", r"\bCrabkey\b"),
    _p("system", "PMX+", r"\bPMX\s?(?:\+|PLUS)"),
    # Codes and identifiers
    _p("template", None, r"\bINT-[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+\b"),
    _p("template", None, r"\bCert_[A-Za-z0-9_-]+"),
    _p("api_test", None, r"\bTC\d{3}_[A-Za-z_]+\b"),
    _p("config_key", None, r"(?<![@\w.])[A-Z][A-Z0-9]{2,}[A-Za-z0-9]*(?:\.[A-Za-z_]+)+", case=True),
    _p("db_table", "PROCEDURECHRG", r"\b(?:medical\.)?Procedurechrg\b"),
    _p("db_table", "DOCUMENT_REPO", r"\bDocument_Repo\b"),
    _p("db_table", "PORTAL_PREF", r"\bPORTAL_PREF\b"),
    _p("hl7_field", None, r"\b(?:MSH|PID|PV1|OBR|OBX|ORC|RXA|RXR|NK1|IN1)[-|]\d+(?:\.\d+)*", case=True),
    _p("loinc", None, r"\b\d{4,5}-\d\b"),
]

PATIENT_RE = re.compile(
    r"\bpat(?:ient|inent|no)?[ \t]*(?:#|no\.?|ids?|number)?[ \t]*[-:]?[ \t]*\(?[ \t]*#?"
    r"(\d{2,6}(?:[ \t]*[,&][ \t]*#?\d{2,6})*)", re.I)
ENV_RE = re.compile(r"\bID\s?(\d{1,3})\b")
LABELED_RE = {
    "medication": re.compile(r"^[ \t]*[-~*•]?[ \t]*(?:Drug Description|Medication)[ \t]*[:\-][ \t]*(.+)$",
                             re.I | re.M),
    "pharmacy": re.compile(r"^[ \t]*[-~*•]?[ \t]*Pharmacy[ \t]*[:\-][ \t]*(.+)$", re.I | re.M),
}

# Sheet-level facts, shown in the report but not listed as rows to check
INFO_FLAGS = {"HEADER_TYPO_FIXED", "TITLE_ONLY"}

FLAG_HELP = {
    "MISSING_TITLE": "No title cell; the title was taken from the expected result or first step.",
    "MISSING_STEPS": "The sheet has a steps column but this row's cell is empty.",
    "MISSING_EXPECTED": "The sheet has an expected-result column but this row's cell is empty.",
    "TITLE_ONLY": "The row has only a title (no steps or expected result).",
    "HEADER_TYPO_FIXED": "A column header had a typo (e.g. 'Excepted Result') and was mapped anyway.",
    "STATUS_COMMENT_CONFLICT": "Status says PASS but the comment mentions pending / fail / issue.",
    "UNKNOWN_STATUS": "Status value not recognised; kept as-is in status_raw.",
    "STATUS_IN_OWNER_COLUMN": "A status word (e.g. 'Fail') was typed in the owner column.",
    "DUPLICATE_TEST": "Another test in the same section has the same title AND the same steps (likely a copy).",
    "DUPLICATE_TITLE": "Another test in the same section has the same title but different steps.",
    "MERGED_CONTINUATION_ROW": "Following row(s) had no title and were merged into this test.",
}


# ---------------------------------------------------------------------------
# 2. SMALL HELPERS
# ---------------------------------------------------------------------------

NUM_RE = re.compile(r"^\d+(?:\.\d+)?$")
STEP_RE = re.compile(r"^\s*(?:step\s*)?(\d{1,2})\s*[.):]\s*(?!\d)", re.I)
GROUP_RE = re.compile(r"^(?:[A-Za-z+]+\s*-\s*)?test\s*case\s*-?\s*\d+$|^EPCS\s*-\s*\d+$", re.I)
CONTEXT_LABEL_RE = re.compile(
    r"^\s*(?P<label>pre[-\s]?conditions?|configurations?|config|(?:kno2\s+)?credentials|topic)\b"
    r"\s*[:\-]?\s*", re.I)
KV_RE = re.compile(r"^\s*[-~*•]?\s*([A-Za-z][A-Za-z /#()]{1,40}?)(?:\s+[-–]\s*|\s*[-–]\s+|\s*:\s*)(.*)$")
BODY_ROLES = ("steps", "expected", "expected_2", "expected_3", "test_data", "preconditions", "notes")


def clean(value):
    """Cell value -> trimmed text ('' for empty). Floats like 1.0 become '1'."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = "\n".join(line.rstrip() for line in text.split("\n")).strip()
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "" if text in {"-", "–", "—"} else text


def first_line(text):
    return text.strip().split("\n", 1)[0].strip()


def rest_lines(text):
    parts = text.strip().split("\n", 1)
    return parts[1].strip() if len(parts) > 1 else ""


def slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_") or "x"


def norm_title(text):
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def dedupe(items):
    return list(dict.fromkeys(i for i in items if i))


def is_run_role(role):
    return bool(role) and role.startswith(("owner:", "status:", "comment:"))


def is_flag_role(role):
    return role in ("high_level", "tc_source") or is_run_role(role)


def fence_if_payload(text):
    """Wrap HL7 / JSON / XML payloads in a code fence so Markdown keeps them intact."""
    if re.search(r"^(?:MSH|PID|OBR|OBX)\|", text, re.M) or text.lstrip().startswith(("{", "<")):
        return f"```\n{text}\n```"
    return text


def header_role(text):
    n = re.sub(r"\s+", " ", text.strip().lower())
    if n.startswith("owner"):
        if "mysql" in n:
            return "owner:MySQL"
        if "fb" in n or "firebird" in n:
            return "owner:Firebird"
        return "owner:default"
    return HEADER_SYNONYMS.get(n)


def detect_header(cells):
    """Return {col: (role, header_text)} when the row looks like a header row."""
    if any("\n" in t or len(t) > 45 for t in cells.values()):
        return None
    found = {}
    for col, text in cells.items():
        role = header_role(text)
        if role:
            found[col] = (role, text)
    if len(found) >= 2 and len(found) >= len(cells) - 1:
        return found
    return None


def label_kind(label):
    label = label.lower()
    if label.startswith("pre"):
        return "PRECONDITION"
    if label.startswith(("config", "kno2", "credential")):
        return "CONFIGURATION"
    return "NOTE"


def split_steps(text):
    """Split numbered steps ('1. ...', 'Step 2: ...') into a list.
    Returns (steps, precondition_lines). Lines before the first number that start
    with 'Precondition' are returned separately."""
    if not text:
        return [], []
    lines = [l for l in text.split("\n") if l.strip()]
    if not any(STEP_RE.match(l) for l in lines):
        return [l.strip(" \t-•*~") for l in lines if l.strip(" \t-•*~")], []
    steps, pre, lead, cur = [], [], [], None
    for line in lines:
        m = STEP_RE.match(line)
        if m:
            if cur is not None:
                steps.append(cur)
            cur = line[m.end():].strip()
        elif cur is None:
            lead.append(line.strip())
        else:
            cur += "\n" + line.strip()
    if cur is not None:
        steps.append(cur)
    # Lines before step 1: 'Precondition ...' lines move out; the rest stay as an intro item
    intro = []
    for line in lead:
        if re.match(r"\(?\s*pre[-\s]?cond", line, re.I):
            pre.append(line.strip("() "))
        else:
            intro.append(line)
    if intro:
        steps.insert(0, "\n".join(intro))
    return steps, pre


def parse_kv(text):
    """'Patient ID - 118' / 'Pharmacy: X' lines -> dict (only lines that look like key/value)."""
    data = {}
    for line in text.split("\n"):
        m = KV_RE.match(line)
        if not m:
            continue
        key, value = m.group(1).strip().lower(), m.group(2).strip()
        if not value:
            continue
        data[KV_KEYS.get(key, slug(key).lower())] = value
    if re.search(r"\bnon[-\s]?controll?(?:er|ed)\b", text, re.I):
        data["substance_class"] = "NON_CONTROLLED"
    elif re.search(r"\bcontroll?(?:er|ed)\s+substances?\b", text, re.I):
        data["substance_class"] = "CONTROLLED"
    return data


def extract_patients(text):
    ids = []
    for m in PATIENT_RE.finditer(text):
        ids += re.findall(r"\d{2,6}", m.group(1))
    return dedupe(ids)


def extract_envs(text):
    envs = []
    for m in ENV_RE.finditer(text):
        before = text[max(0, m.start() - 9):m.start()].lower()
        if not before.endswith(("patient ", "fax ")):
            envs.append(f"ID{m.group(1)}")
    return dedupe(envs)


def env_sort(envs):
    return sorted(set(envs), key=lambda e: int(re.sub(r"\D", "", e) or 0))


def extract_entities(text):
    """Return (entities, aliases). Dictionary/regex matches only - no guessing."""
    found, aliases = {}, set()
    for p in ENTITY_PATTERNS:
        for m in p["re"].finditer(text):
            surface = m.group(0).strip()
            value = p["value"] or surface.rstrip("|")
            key = (p["type"], value)
            if key not in found:
                ent = {"type": p["type"], "value": value}
                if p["standard"]:
                    ent["standard"] = p["standard"]
                found[key] = ent
            if p["type"] == "message_type" and p["value"] and \
                    re.sub(r"\W", "", surface).lower() != re.sub(r"\W", "", value).lower():
                aliases.add(surface)
    for pid in extract_patients(text):
        found.setdefault(("patient", pid), {"type": "patient", "value": pid})
    for etype, rx in LABELED_RE.items():
        for m in rx.finditer(text):
            value = m.group(1).strip()
            if value and len(value) <= 120:
                found.setdefault((etype, value), {"type": etype, "value": value})
    return list(found.values()), sorted(aliases)


# ---------------------------------------------------------------------------
# 3. FEATURES (IntropsFeatures sheet + one per module sheet)
# ---------------------------------------------------------------------------

def feature_id(code, name):
    base = re.sub(r"\(.*", "", name).strip() or name
    return f"FEAT:{code}:{slug(base)}"


def parse_doc_links(text, hyperlink):
    links = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        m = re.search(r"https?://\S+", line)
        if m:
            title = line[:m.start()].strip(" :-–") or None
            links.append({"title": title, "url": m.group(0)})
        else:
            links.append({"title": line, "url": hyperlink})
    return links


def parse_features(ws, code, domain, workbook, ingested_at):
    """IntropsFeatures: # | env IDs | name | action/workflow | inbound | outbound | docs | setup."""
    features, by_number, last = {}, {}, None
    orphan_links = []
    for row in ws.iter_rows(min_row=2):
        cells = {get_column_letter(c.column): clean(c.value) for c in row}
        links = {get_column_letter(c.column): c.hyperlink.target for c in row
                 if getattr(c, "hyperlink", None) is not None}
        num, envs_txt, name = cells.get("A", ""), cells.get("B", ""), cells.get("C", "")
        if not any(cells.values()):
            continue
        env_part, _, env_note = envs_txt.partition("(")
        envs = [f"ID{n}" for n in re.findall(r"\d+", env_part)]
        if name and num:
            fid = feature_id(code, name)
            parent = None
            if "." in num:
                parent = by_number.get(num.split(".")[0])
            doc = {
                "_id": fid, "doc_type": "feature", "schema_version": SCHEMA_VERSION,
                "name": name, "domain": domain, "catalog_no": num, "parent": parent,
                "test_sheet": FEATURE_SHEET_MAP.get(name),
                "environments": env_sort(envs),
                "environment_note": env_note.strip(" )") or None,
                "workflow": cells.get("D") if cells.get("D") not in ("", "Action") else None,
                "directions": dedupe([cells.get("E"), cells.get("F")]),
                "setup": cells.get("H") or None,
                "sub_features": [],
                "doc_links": parse_doc_links(cells.get("G", ""), links.get("G")),
                "source": {"workbook": workbook, "sheet": ws.title, "row": row[0].row,
                           "ingested_at": ingested_at},
                "test_count": 0,
            }
            features[fid] = doc
            by_number[num] = fid
            last = doc
        elif name and last:
            last["sub_features"].append({"code": name, "action": cells.get("D") or None,
                                         "environments": env_sort(envs)})
        elif cells.get("G"):
            orphan_links += parse_doc_links(cells["G"], links.get("G"))
    # Loose document links (e.g. 'KNO2_eFax_Integration') -> feature named in the title
    for link in orphan_links:
        title = (link["title"] or "").lower()
        owner = next((f for f in features.values() if f["name"].lower() in title), None)
        if owner:
            owner["doc_links"].append(link)
    return features


def auto_feature(code, domain, workbook, sheet, ingested_at):
    return {
        "_id": f"FEAT:{code}:{slug(sheet)}", "doc_type": "feature", "schema_version": SCHEMA_VERSION,
        "name": sheet, "domain": domain, "catalog_no": None, "parent": None,
        "test_sheet": sheet, "environments": [], "environment_note": None, "workflow": None,
        "directions": [], "setup": None, "sub_features": [], "doc_links": [],
        "source": {"workbook": workbook, "sheet": sheet, "row": None, "ingested_at": ingested_at},
        "test_count": 0,
    }


def render_feature_md(f):
    lines = [f"# Feature: {f['name']} ({f['domain']})",
             f"**Feature ID:** {f['_id']}" + (f" · **Catalog #:** {f['catalog_no']}" if f["catalog_no"] else ""),
             ""]
    if f["test_sheet"]:
        lines.append(f"**Test sheet:** {f['test_sheet']} · **Test cases:** {f['test_count']}")
    if f["environments"]:
        note = f" ({f['environment_note']})" if f["environment_note"] else ""
        lines.append(f"**Environments:** {', '.join(f['environments'])}{note}")
    if f["parent"]:
        lines.append(f"**Parent feature:** {f['parent']}")
    if f["directions"]:
        lines.append(f"**Directions:** {'; '.join(f['directions'])}")
    if f["workflow"]:
        lines += ["", "### Workflow", f["workflow"]]
    if f["sub_features"]:
        lines += ["", "### Sub-features"]
        lines += [f"- {s['code']}" + (f" — {s['action']}" if s["action"] else "") for s in f["sub_features"]]
    if f["setup"]:
        lines += ["", "### Setup", f["setup"]]
    if f["doc_links"]:
        lines += ["", "### Documents"]
        lines += [f"- [{l['title'] or l['url']}]({l['url']})" if l["url"] else f"- {l['title']}"
                  for l in f["doc_links"]]
    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------
# 4. SHEET PARSER
# ---------------------------------------------------------------------------

class SheetParser:
    """Walks one sheet row by row and sorts each row into:
    header / group (test-data block) / context / section / test / continuation / reference."""

    def __init__(self, ws, cfg, code, domain, workbook, feature, ingested_at):
        self.ws, self.cfg = ws, cfg
        self.code, self.domain, self.workbook = code, domain, workbook
        self.feature, self.ingested_at = feature, ingested_at
        self.sheet_slug = slug(ws.title)
        self.title_only_mode = cfg.get("title_only_rows", "section")
        self.colmap, self.col_headers, self.typo_roles = {}, {}, set()
        self.header_seen, self.header_rows = False, []
        if cfg.get("first_data_row"):
            self.colmap = dict(cfg.get("columns", {}))
        self.group = self.section = None
        # sheet/group/section: where a context applies. "block" = a precondition that
        # applies to the following tests until the next precondition or group.
        self.scoped = {"sheet": [], "group": [], "block": [], "section": []}
        self.tests, self.contexts = [], []
        self.prev_kind = None
        self.stats = Counter()
        self.banners = {(mr.min_row, get_column_letter(mr.min_col))
                        for mr in ws.merged_cells.ranges if mr.max_col > mr.min_col}

    # -- row loop ----------------------------------------------------------
    def run(self):
        for row in self.ws.iter_rows():
            cells = {get_column_letter(c.column): clean(c.value) for c in row}
            cells = {k: v for k, v in cells.items() if v}
            if cells:
                self.stats["rows_read"] += 1
                self.handle_row(row[0].row, cells)
        return self

    def handle_row(self, r, cells):
        first_data_row = self.cfg.get("first_data_row")
        if first_data_row and r >= first_data_row:
            self.header_seen = True
        found = None if first_data_row else detect_header(cells)
        if found:
            self.set_header(r, found)
            return

        first_col, first_text = next(iter(cells.items()))
        if GROUP_RE.match(first_line(first_text)):
            others = {c: t for c, t in cells.items() if c != first_col}
            self.start_group(r, first_line(first_text), others)
            return

        fields, extras, roles = self.map_fields(cells)
        run_info = any(is_run_role(k) for k in fields)

        # Precondition / Configuration rows -> context documents
        label_col = next((c for c, t in cells.items() if not NUM_RE.match(t)), None)
        m = CONTEXT_LABEL_RE.match(cells[label_col]) if label_col else None
        if m and not fields.get("expected") and not run_info:
            label_text = cells[label_col]
            whole = label_text.strip().rstrip(":-– ").lower() == m.group("label").lower()
            others = [t for c, t in cells.items() if c != label_col and not NUM_RE.match(t)]
            body = "\n\n".join(others) if whole else "\n\n".join([label_text] + others)
            if body or not self.header_seen:
                self.add_context(label_kind(m.group("label")), r, body or label_text,
                                 label=first_line(label_text))
            else:
                self.set_section(label_text, r, roles.get(label_col))
            return

        # Rows before the first header row: setup notes or a section name
        if not self.header_seen:
            values = list(cells.values())
            if len(values) == 1 and NUM_RE.match(values[0]):
                self.stats["ignored_rows"] += 1
            elif len(values) == 1 and len(values[0]) <= 60 and "\n" not in values[0]:
                self.set_section(values[0], r, None)
            else:
                multiline = any("\n" in v for v in values)
                text = "\n\n".join(values) if multiline else " | ".join(values)
                kind = "PRECONDITION" if re.search(r"pre[-\s]?condition", text, re.I) else \
                    "CONFIGURATION" if re.search(r"configur", text, re.I) else "NOTE"
                self.add_context(kind, r, text, label=first_line(values[0])[:80], merge=True)
            return

        # Section header rows
        content = {c: t for c, t in cells.items() if not is_flag_role(roles.get(c))}
        if not content:
            self.add_reference(r, cells)
            return
        if len(content) == 1:
            (col, text), = content.items()
            if NUM_RE.match(text):
                self.stats["ignored_rows"] += 1
                return
            role = roles.get(col)
            if (r, col) in self.banners or (
                    len(text) <= 150 and (role != "title" or self.title_only_mode == "section")):
                self.set_section(text, r, role)
                return
        if self.looks_like_described_section(fields, run_info):
            self.set_section(fields["title"] + "\n" + fields["steps"], r, "title")
            return

        # Test rows
        title = fields.get("title")
        body = any(fields.get(k) for k in BODY_ROLES)
        if fields.get("label"):
            self.set_section(fields["label"], r, "label", inline=True)
        if title and (body or run_info or extras or fields.get("tc_number")
                      or self.title_only_mode == "test"):
            self.new_test(r, fields, extras)
        elif title:
            self.set_section(title, r, "title")
        elif body or run_info:
            if not fields.get("tc_number") and not run_info and self.prev_kind == "test":
                self.continue_test(r, fields, extras)
            elif fields.get("tc_number") or run_info or (fields.get("steps") and fields.get("expected")):
                self.new_test(r, fields, extras)
            else:
                self.add_reference(r, cells)
        else:
            self.add_reference(r, cells)

    # -- row types ---------------------------------------------------------
    def set_header(self, r, found):
        colmap, track, self.typo_roles = {}, None, set()
        self.col_headers = {}
        for col in sorted(found, key=column_index_from_string):
            role, text = found[col]
            if role.startswith("owner:"):
                track = role.split(":", 1)[1]
            elif role == "status":
                role = f"status:{track or 'default'}"
            elif role == "comment":
                role = f"comment:{track}" if track else "notes"
            colmap[col] = role
            self.col_headers[col] = text
            if re.sub(r"\s+", " ", text.strip().lower()) in TYPO_HEADERS:
                self.typo_roles.add(role)
        colmap.update(self.cfg.get("columns", {}))
        self.colmap = colmap
        self.header_seen = True
        self.header_rows.append(r)
        self.prev_kind = "header"

    def map_fields(self, cells):
        fields, extras, roles = {}, {}, {}
        for col, text in cells.items():
            role = self.colmap.get(col)
            if role is None and col == "A":
                role = "tc_number" if NUM_RE.match(text) else "label"
            elif role == "tc_number" and not NUM_RE.match(text):
                role = "label"
            roles[col] = role
            if role is None:
                extras[col] = text
            elif role in fields:
                fields[role] += "\n\n" + text
            else:
                fields[role] = text
        return fields, extras, roles

    def header_for(self, col):
        return self.col_headers.get(col) or f"Column {col}"

    def scope(self):
        if self.section and self.header_seen:
            return "section"
        return "group" if self.group else "sheet"

    def start_group(self, r, name, others):
        self.group, self.section = name, None
        self.scoped["group"], self.scoped["block"], self.scoped["section"] = [], [], []
        self.stats["groups"] += 1
        self.prev_kind = "group"
        if others:
            text = "\n\n".join(others.values())
            data = parse_kv(text)
            kind = "TEST_DATA" if len(data) >= 2 else "NOTE"
            self.add_context(kind, r, text, label=name, data=data)

    def set_section(self, text, r, role, inline=False):
        if role == "label" and not inline and self.cfg.get("section_resets_group"):
            self.group = None
            self.scoped["group"] = []
        self.section = first_line(text)[:150]
        self.scoped["section"] = []
        self.stats["sections"] += 1
        self.prev_kind = "section"
        detail = rest_lines(text)
        if detail:
            self.add_context("NOTE", r, detail, label=self.section)

    def add_context(self, kind, r, text, label=None, data=None, merge=False):
        scope = self.scope()
        if scope == "section" and kind in ("PRECONDITION", "CONFIGURATION"):
            scope = "block"
            self.scoped["block"] = []
        last = self.contexts[-1] if self.contexts else None
        if merge and last and self.prev_kind == "context" and last["scope"] == scope:
            last["text"] += ("\n" if " | " in text else "\n\n") + text
            last["rows"].append(r)
            return
        rec = {"id": f"CTX:{self.code}:{self.sheet_slug}:R{r}", "kind": kind, "rows": [r],
               "scope": scope, "group": self.group, "section": self.section,
               "label": label, "text": text, "data": data or {}}
        self.contexts.append(rec)
        self.scoped[scope].append(rec["id"])
        self.stats["contexts"] += 1
        self.prev_kind = "context"

    def add_reference(self, r, cells):
        text = "\n".join(f"[{self.header_for(c)}] {t}" if len(cells) > 1 else t
                         for c, t in cells.items())
        last = self.contexts[-1] if self.contexts else None
        self.stats["reference_rows"] += 1
        if last and self.prev_kind == "reference" and last["kind"] == "REFERENCE":
            last["text"] += "\n\n" + text
            last["rows"].append(r)
            return
        self.add_context("REFERENCE", r, text, label=self.section)
        self.prev_kind = "reference"

    def looks_like_described_section(self, fields, run_info):
        """ALL MODULES: 'Nurse Only Visit' + a one-line description, no owner, no numbered steps."""
        has_owner_cols = any(r.startswith("owner:") for r in self.colmap.values())
        return (has_owner_cols and not run_info and set(fields) == {"title", "steps"}
                and len(fields["title"]) <= 80
                and not any(STEP_RE.match(l) for l in fields["steps"].split("\n")))

    def new_test(self, r, fields, extras):
        flags = ["HEADER_TYPO_FIXED"] if self.typo_roles & set(fields) else []
        self.tests.append({
            "row": r, "rows": [r], "fields": dict(fields), "extras": dict(extras),
            "extra_headers": {c: self.header_for(c) for c in extras},
            "group": self.group, "section": self.section,
            "ctx_ids": (self.scoped["sheet"] + self.scoped["group"] + self.scoped["block"]
                        + self.scoped["section"]),
            "run_roles": [role for col, role in sorted(self.colmap.items(),
                                                      key=lambda kv: column_index_from_string(kv[0]))
                          if is_run_role(role)],
            "layout_roles": set(self.colmap.values()),
            "role_headers": {role: self.header_for(col) for col, role in self.colmap.items()},
            "flags": flags,
        })
        self.stats["tests"] += 1
        self.prev_kind = "test"

    def continue_test(self, r, fields, extras):
        rec = self.tests[-1]
        for role, text in fields.items():
            rec["fields"][role] = (rec["fields"][role] + "\n\n" + text) if rec["fields"].get(role) else text
        for col, text in extras.items():
            rec["extras"][col] = (rec["extras"][col] + "\n\n" + text) if rec["extras"].get(col) else text
            rec["extra_headers"].setdefault(col, self.header_for(col))
        rec["rows"].append(r)
        if "MERGED_CONTINUATION_ROW" not in rec["flags"]:
            rec["flags"].append("MERGED_CONTINUATION_ROW")
        self.stats["continuation_rows"] += 1


# ---------------------------------------------------------------------------
# 5. DOCUMENT BUILDERS
# ---------------------------------------------------------------------------

def normalize_status(raw, flags):
    if not raw:
        return "NOT_RUN"
    status = STATUS_MAP.get(raw.strip().lower())
    if status is None:
        flags.append("UNKNOWN_STATUS")
        return "UNKNOWN"
    return status


def build_runs(rec, ingested_at, flags):
    f, tracks = rec["fields"], {}
    for role in rec["run_roles"]:
        kind, track = role.split(":", 1)
        tracks.setdefault(track, False)
        if kind == "owner":
            tracks[track] = True
    runs = []
    for track, has_owner_col in tracks.items():
        owner = f.get(f"owner:{track}")
        status_raw = f.get(f"status:{track}")
        comment = f.get(f"comment:{track}")
        if has_owner_col and not (owner or status_raw or comment):
            continue  # this test is not assigned to this track
        if owner and owner.strip().lower() in STATUS_MAP and owner.strip().lower() not in ("n/a", "na"):
            flags.append("STATUS_IN_OWNER_COLUMN")
            status_raw, owner = status_raw or owner, None
        status = normalize_status(status_raw, flags)
        if owner and owner.strip().lower() in ("n/a", "na"):
            owner = None
            if status == "NOT_RUN":
                status = "NA"
        if status == "PASS" and comment and CONFLICT_COMMENT_RE.search(comment):
            flags.append("STATUS_COMMENT_CONFLICT")
        runs.append({"track": track, "scope": TRACK_SCOPES.get(track), "owner": owner,
                     "status": status, "status_raw": status_raw or None, "comment": comment or None,
                     "executed_at": None, "recorded_at": ingested_at})
    return runs


def breadcrumb(domain, sheet, group, section):
    return " > ".join(x for x in (domain, sheet, group, section) if x)


def derive_title(f):
    for key in ("expected", "steps", "test_data", "notes"):
        if f.get(key):
            line = STEP_RE.sub("", first_line(f[key])).strip()
            return line[:117] + "..." if len(line) > 120 else line
    return "(untitled)"


def render_test_md(doc, rec, expected_parts):
    f, src = rec["fields"], doc["source"]
    rows = ", ".join(str(r) for r in rec["rows"])
    meta = [f"**Test ID:** {doc['_id']}"]
    if doc["test_number"]:
        meta.append(f"**Test #:** {doc['test_number']}")
    meta.append(f"**Source:** {src['workbook']} › {src['sheet']} › row {rows}")
    lines = [f"# {doc['hierarchy']['breadcrumb']}", f"## {doc['title']}", "", " · ".join(meta) + "  "]
    if doc["context_snapshot"]:
        lines.append("**Context:** " + "; ".join(f"{k.replace('_', ' ')}: {v}"
                                                   for k, v in doc["context_snapshot"].items()) + "  ")
    if doc["environments"]:
        lines.append("**Environments:** " + ", ".join(doc["environments"]))
    lines.append("")

    def block(heading, text):
        if text:
            lines.extend([f"### {heading}", fence_if_payload(text), ""])

    block("Preconditions", f.get("preconditions"))
    block("Test data", f.get("test_data"))
    block("Steps", f.get("steps"))
    if len(expected_parts) > 1:
        for header, text in expected_parts:
            block(header, text)
    elif expected_parts:
        block("Expected result", expected_parts[0][1])
    block("Notes", f.get("notes"))

    other = []
    if f.get("tc_source"):
        other.append(("TC source", f["tc_source"]))
    if f.get("high_level"):
        other.append(("HighLevel", f["high_level"]))
    if f.get("label") and "\n" in f["label"]:
        other.append(("Section label", f["label"].replace("\n", " ")))
    for col, text in rec["extras"].items():
        other.append((rec["extra_headers"].get(col, f"Column {col}"), text.replace("\n", " / ")))
    if other:
        lines.append("### Other columns")
        lines += [f"- **{h}:** {v}" for h, v in other]
        lines.append("")
    if doc["runs"]:
        lines += ["### Execution", "| Track | Scope | Owner | Status | Comment |", "|---|---|---|---|---|"]
        for run in doc["runs"]:
            lines.append("| " + " | ".join(str(run[k] or "-").replace("|", "/").replace("\n", " ")
                                            for k in ("track", "scope", "owner", "status_raw", "comment")) + " |")
    return "\n".join(lines).strip() + "\n"


def build_test_doc(p, rec, seq, ctx_by_id):
    f, flags = rec["fields"], list(rec["flags"])
    tc_id = f"TC:{p.code}:{p.sheet_slug}:R{rec['row']}"
    title = f.get("title")
    if not title:
        title = derive_title(f)
        flags.append("MISSING_TITLE")
    title = re.sub(r"\s+", " ", title).strip()

    steps, pre_from_steps = split_steps(f.get("steps", ""))
    preconditions = dedupe(([f["preconditions"]] if f.get("preconditions") else []) + pre_from_steps)
    expected_parts = [(rec["role_headers"].get(role, role), f[role])
                      for role in ("expected", "expected_2", "expected_3") if f.get(role)]
    expected = "\n\n".join(f"{h}: {t}" for h, t in expected_parts) if len(expected_parts) > 1 else \
        (expected_parts[0][1] if expected_parts else None)

    missing = []
    if "steps" in rec["layout_roles"] and not f.get("steps"):
        missing.append("steps")
    if "expected" in rec["layout_roles"] and not expected:
        missing.append("expected_result")
    if not f.get("steps") and not expected and not f.get("test_data"):
        flags.append("TITLE_ONLY")
    else:
        flags += [f"MISSING_{m.split('_')[0].upper()}" for m in missing]

    contexts = [ctx_by_id[c] for c in rec["ctx_ids"] if c in ctx_by_id]
    snapshot = {}
    for ctx in contexts:
        if ctx["kind"] == "TEST_DATA":
            snapshot.update({k: ctx["data"][k] for k in SNAPSHOT_KEYS if k in ctx["data"]})

    own_text = "\n".join([title] + [t for k, t in f.items() if not is_run_role(k)] + list(rec["extras"].values()))
    entities, aliases = extract_entities(own_text)
    seen = {(e["type"], e["value"]) for e in entities}
    for ctx in contexts:
        ctx_ents, _ = extract_entities(ctx["text"])
        for k, v in ctx["data"].items():
            if k in ("patient_id", "medication", "pharmacy"):
                ctx_ents.append({"type": "patient" if k == "patient_id" else k, "value": v})
        for e in ctx_ents:
            if e["type"] in ("patient", "medication", "pharmacy") and (e["type"], e["value"]) not in seen:
                seen.add((e["type"], e["value"]))
                entities.append({**e, "from": "context"})

    environments = env_sort(p.feature["environments"] + extract_envs(own_text)
                            + [env for ctx in contexts for env in extract_envs(ctx["text"])])
    runs = build_runs(rec, p.ingested_at, flags)
    attributes = {}
    if f.get("tc_source"):
        attributes["tc_source"] = f["tc_source"]
    if f.get("high_level"):
        attributes["is_high_level"] = f["high_level"].strip().lower() == "yes"
    if rec["extras"]:
        attributes["extra_columns"] = {rec["extra_headers"].get(c, c): t for c, t in rec["extras"].items()}

    doc = {
        "_id": tc_id,
        "doc_type": "test_case",
        "schema_version": SCHEMA_VERSION,
        "source": {
            "workbook": p.workbook, "sheet": p.ws.title, "row": rec["row"],
            **({"rows": rec["rows"]} if len(rec["rows"]) > 1 else {}),
            "raw_headers": {r: h for r, h in rec["role_headers"].items() if r in f},
            "content_hash": None, "ingested_at": p.ingested_at,
        },
        "hierarchy": {
            "domain": p.domain, "module": p.ws.title, "group": rec["group"], "section": rec["section"],
            "sequence_no": seq, "breadcrumb": breadcrumb(p.domain, p.ws.title, rec["group"], rec["section"]),
        },
        "test_number": f.get("tc_number"),
        "title": title,
        "preconditions": preconditions,
        "steps": steps,
        "expected_result": expected,
        "test_data": f.get("test_data"),
        "notes": f.get("notes"),
        "attributes": attributes,
        "context_refs": [c["id"] for c in contexts],
        "context_snapshot": snapshot,
        "entities": entities,
        "environments": environments,
        "runs": runs,
        "relations": [{"type": "BELONGS_TO", "target": p.feature["_id"]}]
                     + [{"type": "USES_CONTEXT", "target": c["id"]} for c in contexts],
        "quality": {"flags": sorted(set(flags)), "missing_fields": missing},
        "search": {
            "keywords": dedupe([f.get("tc_number"), rec["group"], rec["section"]]
                               + [e["value"] for e in entities if e["type"] != "system"] + environments),
            "aliases": aliases,
        },
        "content_md": None,
    }
    doc["content_md"] = render_test_md(doc, rec, expected_parts)
    doc["source"]["content_hash"] = "sha256:" + hashlib.sha256(doc["content_md"].encode("utf-8")).hexdigest()
    return doc


CONTEXT_TITLES = {"TEST_DATA": "Test data", "PRECONDITION": "Precondition", "CONFIGURATION": "Configuration",
                  "NOTE": "Note", "REFERENCE": "Reference data"}


def build_context_doc(p, ctx, test_ids):
    entities, _ = extract_entities(ctx["text"])
    environments = env_sort(extract_envs(ctx["text"]))
    crumb = breadcrumb(p.domain, p.ws.title, ctx["group"],
                       ctx["section"] if ctx["scope"] == "section" or ctx["kind"] == "REFERENCE" else None)
    heading = CONTEXT_TITLES[ctx["kind"]] + (f": {ctx['label']}" if ctx["label"] else "")
    lines = [f"# {crumb}", f"## {heading}", "",
             f"**Context ID:** {ctx['id']} · **Source:** {p.workbook} › {p.ws.title} › row "
             + ", ".join(str(r) for r in ctx["rows"]) + f" · **Applies to:** {ctx['scope']} ({len(test_ids)} tests)", ""]
    if ctx["data"]:
        lines += [f"- **{k.replace('_', ' ')}:** {v}" for k, v in ctx["data"].items()] + [""]
    lines.append(fence_if_payload(ctx["text"]))
    return {
        "_id": ctx["id"],
        "doc_type": "context",
        "schema_version": SCHEMA_VERSION,
        "context_kind": ctx["kind"],
        "label": ctx["label"],
        "source": {"workbook": p.workbook, "sheet": p.ws.title, "rows": ctx["rows"],
                   "ingested_at": p.ingested_at},
        "hierarchy": {"domain": p.domain, "module": p.ws.title, "group": ctx["group"],
                      "section": ctx["section"], "breadcrumb": crumb},
        "applies_to": {"scope": ctx["scope"], "test_ids": test_ids},
        "data": ctx["data"],
        "text": ctx["text"],
        "entities": entities,
        "environments": environments,
        "relations": [{"type": "BELONGS_TO", "target": p.feature["_id"]}],
        "search": {"keywords": dedupe([ctx["label"], ctx["group"]]
                                      + [e["value"] for e in entities if e["type"] != "system"]
                                      + environments)},
        "content_md": "\n".join(lines).strip() + "\n",
    }


def link_same_titles(tests):
    """Same title in another group/section -> SAME_SCENARIO link.
    Same title in the same group and section -> DUPLICATE_TEST (same steps) or DUPLICATE_TITLE."""
    def where(t):
        return t["hierarchy"]["group"], t["hierarchy"]["section"]
    by_title = defaultdict(list)
    for t in tests:
        by_title[norm_title(t["title"])].append(t)
    for same in by_title.values():
        if len(same) < 2:
            continue
        for t in same:
            others = [o for o in same if o is not t]
            twins = [o for o in others if where(o) == where(t)]
            if twins:
                exact = any(o["steps"] == t["steps"] and o["expected_result"] == t["expected_result"] for o in twins)
                flag = "DUPLICATE_TEST" if exact else "DUPLICATE_TITLE"
                t["quality"]["flags"] = sorted(set(t["quality"]["flags"]) | {flag})
            for o in [o for o in others if where(o) != where(t)][:10]:
                t["relations"].append({"type": "SAME_SCENARIO", "target": o["_id"]})


def reference_sheet_context(ws, code, domain, workbook, feature, ingested_at):
    """Sheets that are lookup tables, not tests (e.g. Fhir_appointment_updated)."""
    rows, lines = [], []
    for row in ws.iter_rows():
        cells = [clean(c.value) for c in row]
        cells = [c for c in cells if c]
        if cells:
            rows.append(row[0].row)
            lines.append(" | ".join(cells))
    text = "\n".join(lines)
    cid = f"CTX:{code}:{slug(ws.title)}:R{rows[0] if rows else 1}"
    crumb = breadcrumb(domain, ws.title, None, None)
    return {
        "_id": cid, "doc_type": "context", "schema_version": SCHEMA_VERSION,
        "context_kind": "REFERENCE", "label": ws.title,
        "source": {"workbook": workbook, "sheet": ws.title, "rows": rows, "ingested_at": ingested_at},
        "hierarchy": {"domain": domain, "module": ws.title, "group": None, "section": None, "breadcrumb": crumb},
        "applies_to": {"scope": "sheet", "test_ids": []},
        "data": {}, "text": text,
        "entities": extract_entities(text)[0], "environments": env_sort(extract_envs(text)),
        "relations": [{"type": "BELONGS_TO", "target": feature["_id"]}],
        "search": {"keywords": [ws.title]},
        "content_md": f"# {crumb}\n## Reference data: {ws.title}\n\n**Context ID:** {cid} · "
                      f"**Source:** {workbook} › {ws.title} › {len(rows)} rows\n\n{text}\n",
    }


# ---------------------------------------------------------------------------
# 6. OUTPUT SHAPE
#    The builders above keep working fields (steps list, quality flags, ...) for
#    the duplicate check and the review report. Only the fields below are written.
#    Every document type shares the same core fields:
#      _id, doc_type, schema_version, title, source, hierarchy,
#      entities, environments, relations, content_md, embedding
# ---------------------------------------------------------------------------

HIERARCHY_KEYS = ("domain", "module", "group", "section", "breadcrumb")


def embedding_stub(content_md):
    """Empty embedding slot, filled later from content_md.
    text_hash tells the embedding step which documents changed since the last run."""
    return {"model": None, "dims": None,
            "text_hash": "sha256:" + hashlib.sha256(content_md.encode("utf-8")).hexdigest(),
            "vector": None}


def to_output_test(d):
    return {
        "_id": d["_id"], "doc_type": "test_case", "schema_version": SCHEMA_VERSION,
        "title": d["title"],
        "test_number": d["test_number"],
        "source": {k: d["source"][k] for k in ("workbook", "sheet", "row", "rows", "ingested_at")
                   if k in d["source"]},
        "hierarchy": {k: d["hierarchy"][k] for k in HIERARCHY_KEYS},
        "runs": [{k: v for k, v in r.items() if k != "status_raw"} for r in d["runs"]],
        "entities": d["entities"],
        "environments": d["environments"],
        "relations": d["relations"],
        "content_md": d["content_md"],
        "embedding": embedding_stub(d["content_md"]),
    }


def to_output_context(d):
    kind_title = CONTEXT_TITLES[d["context_kind"]]
    return {
        "_id": d["_id"], "doc_type": "context", "schema_version": SCHEMA_VERSION,
        "title": f"{kind_title}: {d['label']}" if d["label"] else kind_title,
        "context_kind": d["context_kind"],
        "scope": d["applies_to"]["scope"],
        "source": d["source"],
        "hierarchy": {k: d["hierarchy"][k] for k in HIERARCHY_KEYS},
        "entities": d["entities"],
        "environments": d["environments"],
        "relations": d["relations"],
        "content_md": d["content_md"],
        "embedding": embedding_stub(d["content_md"]),
    }


def to_output_feature(d):
    return {
        "_id": d["_id"], "doc_type": "feature", "schema_version": SCHEMA_VERSION,
        "title": d["name"],
        "test_count": d["test_count"],
        "source": d["source"],
        "hierarchy": {"domain": d["domain"], "module": d["test_sheet"], "group": None, "section": None,
                      "breadcrumb": f"{d['domain']} > {d['name']}"},
        "entities": extract_entities(d["content_md"])[0],
        "environments": d["environments"],
        "relations": [{"type": "PART_OF", "target": d["parent"]}] if d["parent"] else [],
        "content_md": d["content_md"],
        "embedding": embedding_stub(d["content_md"]),
    }


# ---------------------------------------------------------------------------
# 7. MAIN
# ---------------------------------------------------------------------------

def parse_workbook(path, ingested_at, report):
    workbook = path.stem
    code, domain = WORKBOOKS.get(workbook, (slug(workbook).upper()[:12], workbook))
    wb = openpyxl.load_workbook(path, data_only=True)
    features = {}
    if "IntropsFeatures" in wb.sheetnames:
        features = parse_features(wb["IntropsFeatures"], code, domain, workbook, ingested_at)
    tests, contexts = [], []
    for ws in wb.worksheets:
        cfg = SHEET_CONFIG.get(ws.title, {})
        mode = cfg.get("mode", "tests")
        if mode in ("skip", "features"):
            report["sheets"].append({"workbook": workbook, "sheet": ws.title, "mode": mode})
            continue
        feature = next((f for f in features.values() if f["test_sheet"] == ws.title), None)
        if feature is None:
            feature = auto_feature(code, domain, workbook, ws.title, ingested_at)
            features[feature["_id"]] = feature
        if mode == "reference":
            contexts.append(reference_sheet_context(ws, code, domain, workbook, feature, ingested_at))
            report["sheets"].append({"workbook": workbook, "sheet": ws.title, "mode": mode, "contexts": 1})
            continue

        p = SheetParser(ws, cfg, code, domain, workbook, feature, ingested_at).run()
        ctx_by_id = {c["id"]: c for c in p.contexts}
        sheet_tests = [build_test_doc(p, rec, i + 1, ctx_by_id) for i, rec in enumerate(p.tests)]
        link_same_titles(sheet_tests)
        applies = defaultdict(list)
        for t in sheet_tests:
            for cid in t["context_refs"]:
                applies[cid].append(t["_id"])
        contexts += [build_context_doc(p, c, applies[c["id"]]) for c in p.contexts]
        tests += sheet_tests
        feature["test_count"] += len(sheet_tests)
        report["sheets"].append({
            "workbook": workbook, "sheet": ws.title, "mode": mode, "stats": dict(p.stats),
            "header_rows": p.header_rows,
            "columns": {col: (p.col_headers.get(col), role) for col, role in sorted(
                p.colmap.items(), key=lambda kv: column_index_from_string(kv[0]))},
            "flags": Counter(fl for t in sheet_tests for fl in t["quality"]["flags"]),
            "flagged": [(t["source"]["row"], t["_id"],
                         [f for f in t["quality"]["flags"] if f not in INFO_FLAGS], t["title"])
                        for t in sheet_tests if set(t["quality"]["flags"]) - INFO_FLAGS],
            "tests": len(sheet_tests),
        })
    for f in features.values():
        f["content_md"] = render_feature_md(f)
    return tests, contexts, list(features.values())


def write_report(path, report, totals, inputs):
    L = ["# Parser review report", "",
         f"Generated: {report['generated']}  ",
         f"Inputs: {', '.join(inputs)}", "",
         "## Totals", "", "| Document type | Count |", "|---|---|"]
    L += [f"| {k} | {v} |" for k, v in totals.items()]
    L += ["", "## Per sheet", "",
          "| Workbook | Sheet | Mode | Rows read | Tests | Sections | Contexts | Reference rows | Merged rows | Flagged tests |",
          "|---|---|---|---|---|---|---|---|---|---|"]
    for s in report["sheets"]:
        st = s.get("stats", {})
        L.append(f"| {s['workbook']} | {s['sheet']} | {s['mode']} | {st.get('rows_read', '-')} | "
                 f"{s.get('tests', '-')} | {st.get('sections', '-')} | {st.get('contexts', s.get('contexts', '-'))} | "
                 f"{st.get('reference_rows', '-')} | {st.get('continuation_rows', '-')} | "
                 f"{len(s.get('flagged', [])) if 'flagged' in s else '-'} |")
    all_flags = Counter()
    for s in report["sheets"]:
        all_flags.update(s.get("flags", {}))
    L += ["", "## Flags", "", "| Flag | Tests | Meaning |", "|---|---|---|"]
    L += [f"| {k} | {v} | {FLAG_HELP.get(k, '')}{' (info only)' if k in INFO_FLAGS else ''} |"
          for k, v in all_flags.most_common()]
    L += ["", "## Column mapping per sheet", "",
          "Check that each column got the right role. `label` = section name in column A.", ""]
    for s in report["sheets"]:
        if s.get("columns"):
            cols = "; ".join(f"{c}={h or '(no header)'}→{r}" for c, (h, r) in s["columns"].items())
            L.append(f"- **{s['sheet']}** (header row {', '.join(map(str, s['header_rows'])) or 'none'}): {cols}")
    L += ["", "## Rows to check", "",
          "Up to 40 tests per flag per sheet (info-only flags are left out).",
          "STATUS_*, MISSING_TITLE and MERGED_CONTINUATION_ROW are the most worth a look.", ""]
    for s in report["sheets"]:
        if not s.get("flagged"):
            continue
        L.append(f"### {s['sheet']}")
        per_flag = defaultdict(list)
        for row, tid, flags, title in s["flagged"]:
            for fl in flags:
                per_flag[fl].append((row, tid, title))
        for fl, items in sorted(per_flag.items()):
            L.append(f"**{fl}** ({len(items)})")
            L += [f"- R{row} · `{tid}` · {title[:90]}" for row, tid, title in items[:40]]
            if len(items) > 40:
                L.append(f"- ... and {len(items) - 40} more")
            L.append("")
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    here = Path(__file__).resolve().parent
    default_inputs = [str(here / f"{name}.xlsx") for name in WORKBOOKS if (here / f"{name}.xlsx").exists()]
    ap = argparse.ArgumentParser(description="Parse QA test-case Excel files into canonical JSON for MongoDB.")
    ap.add_argument("--input", nargs="+", default=default_inputs, help="Excel files to parse")
    ap.add_argument("--out", default=str(here / "output"), help="Output folder")
    args = ap.parse_args()
    if not args.input:
        sys.exit("No input files found. Pass them with --input.")

    ingested_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    report = {"generated": ingested_at, "sheets": []}
    tests, contexts, features = [], [], []
    for name in args.input:
        path = Path(name)
        print(f"Reading {path.name} ...")
        t, c, f = parse_workbook(path, ingested_at, report)
        tests += t
        contexts += c
        features += f

    tests = [to_output_test(d) for d in tests]
    contexts = [to_output_context(d) for d in contexts]
    features = [to_output_feature(d) for d in features]

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    for fname, docs in (("test_cases.json", tests), ("contexts.json", contexts), ("features.json", features)):
        (out / fname).write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")
    with open(out / "qa_knowledge.jsonl", "w", encoding="utf-8") as fh:
        for doc in features + contexts + tests:
            fh.write(json.dumps(doc, ensure_ascii=False) + "\n")
    totals = {"test_case": len(tests), "context": len(contexts), "feature": len(features)}
    write_report(out / "review_report.md", report, totals, [Path(i).name for i in args.input])

    print(f"\nDone. {totals['test_case']} test cases, {totals['context']} contexts, {totals['feature']} features")
    print(f"Output folder: {out}")
    print("Load into MongoDB with:")
    print(f'  mongoimport --uri "<your-connection-string>" --collection qa_knowledge '
          f'--file "{out / "qa_knowledge.jsonl"}" --mode upsert')


if __name__ == "__main__":
    main()
